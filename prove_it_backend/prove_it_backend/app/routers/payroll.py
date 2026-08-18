import io
import re
from datetime import date, datetime
from typing import Optional

import openpyxl
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from pydantic import BaseModel
from pymongo.database import Database

from app.core import collections, pay_register_schema
from app.core.audit import log_action
from app.core.database import get_db
from app.core.mongo_utils import get_or_404
from app.core.security import get_current_user, require_permission

router = APIRouter()

_BANK_STATUTORY_KEYS = [c["key"] for c in pay_register_schema.BANK_STATUTORY_FIELDS]
_JOB_DETAILS_KEYS = [c["key"] for c in pay_register_schema.JOB_DETAILS_FIELDS]
_EMPLOYEE_DETAILS_EXTRA_KEYS = [c["key"] for c in pay_register_schema.EMPLOYEE_DETAILS_EXTRA_FIELDS]

# designation/department/phone (Job Details) and email (Employee Details) are the exact
# same field the Employees page already manages — shown on these tabs for full-fidelity
# display (see pay_register_schema's docstring), but deliberately absent from the
# corresponding *Update models below so Payroll:edit can't double as a second write path
# for data Employees:edit already owns.


class BankStatutoryUpdate(BaseModel):
    pay_mode: Optional[str] = None
    bank_name: Optional[str] = None
    bank_branch: Optional[str] = None
    ifsc: Optional[str] = None
    bank_ref_no: Optional[str] = None
    cheque_no: Optional[str] = None
    cheque_date: Optional[str] = None
    account_no: Optional[str] = None
    name_as_per_account: Optional[str] = None
    aadhaar: Optional[str] = None
    pf_no: Optional[str] = None
    pf_date: Optional[str] = None
    uan: Optional[str] = None
    esi_no: Optional[str] = None
    esi_date: Optional[str] = None
    esi_office: Optional[str] = None


class JobDetailsUpdate(BaseModel):
    branch_code: Optional[str] = None
    branch: Optional[str] = None
    category: Optional[str] = None
    scale: Optional[str] = None
    shift: Optional[str] = None
    work_location: Optional[str] = None
    mobile: Optional[str] = None
    internal_id: Optional[str] = None
    address_permanent: Optional[str] = None
    address_correspondence: Optional[str] = None


class EmployeeDetailsUpdate(BaseModel):
    dor: Optional[str] = None
    dob: Optional[str] = None
    notice_period: Optional[str] = None
    pan: Optional[str] = None


class CompanyInfoUpdate(BaseModel):
    company_name: Optional[str] = None
    company_address: Optional[str] = None


def _out_employee_fields(e: dict, keys) -> dict:
    out = {"emp_id": e["emp_id"], "name": e["name"]}
    for key in keys:
        out[key] = e.get(key)
    return out


def _normalize(value) -> str:
    return " ".join(str(value).split())


def _find_header_anchor(ws):
    """The Pay Register title/section rows above the real header vary in row count
    between template exports, so anchor on the 'Code' column header instead of a
    fixed row number. Returns (header_row, code_column) or None."""
    for r in range(1, min(ws.max_row, 10) + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and _normalize(v).lower() == "code":
                return r, c
    return None


def _build_column_map(code_col: int) -> dict:
    """Maps schema keys to sheet columns by POSITION relative to the 'Code' column,
    not by header text — several headers repeat verbatim across sections (e.g.
    Earnings' "Gross Earning" vs Company Contribution's "GROSS EARNING"), so a
    text->key lookup would collide and silently drop one of them. SHEET_COLUMNS (the
    template's full physical layout, including Bank & Statutory Details — see
    pay_register_schema's docstring for why that section is excluded further downstream
    instead of here) is defined in the exact left-to-right order the template exports
    them in, with 'code' at index 1 (right after 's_no'), so every other column's offset
    from code_col follows from its index in that list."""
    code_index = next(i for i, c in enumerate(pay_register_schema.SHEET_COLUMNS) if c["key"] == "code")
    start_col = code_col - code_index
    return {c["key"]: start_col + i for i, c in enumerate(pay_register_schema.SHEET_COLUMNS)}


_MONTH_ABBR = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
    "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
}


def _detect_period(ws) -> Optional[str]:
    """Reads the 'For the Month of <Month>/<Year>' title text near the top of the
    sheet and turns it into a sortable 'YYYY-MM' period key."""
    blob = " ".join(
        ws.cell(row=r, column=c).value
        for r in range(1, 5) for c in range(1, ws.max_column + 1)
        if isinstance(ws.cell(row=r, column=c).value, str)
    )
    m = re.search(r"for the month of\s+([A-Za-z]+)\s*[/\-]?\s*(\d{4})", blob, re.IGNORECASE)
    if not m:
        return None
    month_num = _MONTH_ABBR.get(m.group(1)[:3].lower())
    if not month_num:
        return None
    return f"{m.group(2)}-{month_num:02d}"


def _detect_company_info(ws):
    """The template's own title block (e.g. row 2) carries the company name and address
    as one cell, name and address separated by a literal newline — used for the Payslip
    header so it isn't hardcoded to one org. Returns (company_name, company_address),
    either of which may be None if the title block doesn't have this cell (upload still
    succeeds either way; routers/payroll.py's PATCH /{period}/company-info lets it be
    filled in by hand instead). The address's own internal newline (e.g. street line vs
    city/state line) is preserved rather than flattened, so payslip_pdf.py can render it
    as separate lines matching the org's reference template."""
    for r in range(1, 5):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and "\n" in v and "for the month of" not in v.lower():
                name, _, address = v.partition("\n")
                address_lines = [_normalize(line) for line in address.split("\n")]
                address_lines = [line for line in address_lines if line]
                return _normalize(name) or None, ("\n".join(address_lines) or None)
    return None, None


def _cell_text(v):
    if v is None:
        return None
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    if isinstance(v, str):
        return v.strip() or None
    return v


def _cell_number(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, str):
        s = v.strip().replace(",", "")
        if not s:
            return None
        try:
            f = float(s)
            return int(f) if f.is_integer() else f
        except ValueError:
            return v
    return v


def _parse_rows(ws, header_row: int, col_map: dict):
    """Only reads the keys in pay_register_schema.COLUMNS — col_map itself has an entry
    for every SHEET_COLUMNS key (Bank & Statutory Details included), but that section is
    deliberately never stored on the Pay Register (see pay_register_schema's docstring)."""
    name_col = col_map["name"]
    register_keys = [c["key"] for c in pay_register_schema.COLUMNS]
    employees = []
    totals = {}
    for r in range(header_row + 1, ws.max_row + 1):
        name_val = _cell_text(ws.cell(row=r, column=name_col).value)
        if not name_val:
            continue  # blank template row, no employee on it
        if name_val.lower() == "total":
            for key in register_keys:
                if key in pay_register_schema.NUMERIC_KEYS:
                    totals[key] = _cell_number(ws.cell(row=r, column=col_map[key]).value)
            break
        row_data = {}
        for key in register_keys:
            v = ws.cell(row=r, column=col_map[key]).value
            row_data[key] = _cell_number(v) if key in pay_register_schema.NUMERIC_KEYS else _cell_text(v)
        employees.append(row_data)
    return employees, totals


def _out(doc: dict) -> dict:
    return {
        "period": doc["period"],
        "uploaded_at": doc["uploaded_at"],
        "uploaded_by": doc["uploaded_by"],
        "source_filename": doc["source_filename"],
        "employee_count": doc.get("employee_count", len(doc.get("employees", []))),
        "employees": doc.get("employees", []),
        "totals": doc.get("totals", {}),
        "company_name": doc.get("company_name"),
        "company_address": doc.get("company_address"),
    }


def _out_meta(doc: dict) -> dict:
    out = _out(doc)
    out.pop("employees")
    out.pop("totals")
    return out


@router.post("/upload", dependencies=[Depends(require_permission("Payroll", "edit"))])
async def upload(
    file: UploadFile = File(...),
    period: Optional[str] = Form(None),
    db: Database = Depends(get_db),
    cu=Depends(get_current_user),
):
    if not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(400, "Please upload an Excel (.xlsx) file")

    data = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception:
        raise HTTPException(400, "Could not read this file — is it a valid Excel workbook?")
    ws = wb.active

    anchor = _find_header_anchor(ws)
    if anchor is None:
        raise HTTPException(400, "Could not find the Pay Register header row (expected a 'Code' column) in this file")
    header_row, code_col = anchor

    col_map = _build_column_map(code_col)

    detected_period = period or _detect_period(ws)
    if not detected_period:
        raise HTTPException(400, "Could not detect the pay period (month/year) from this file's title row")

    employees, totals = _parse_rows(ws, header_row, col_map)
    if not employees:
        raise HTTPException(400, "No employee rows were found in this file")

    # Preserve a previously-set/manually-entered company name/address across re-uploads
    # if this file's title block doesn't carry one — see PATCH /{period}/company-info.
    existing = db[collections.PAYROLL_REGISTERS].find_one(
        {"org_id": cu.org_id, "period": detected_period}, {"company_name": 1, "company_address": 1},
    )
    detected_name, detected_address = _detect_company_info(ws)
    company_name = detected_name or (existing.get("company_name") if existing else None)
    company_address = detected_address or (existing.get("company_address") if existing else None)

    doc = {
        "org_id": cu.org_id,
        "period": detected_period,
        "uploaded_at": datetime.utcnow().isoformat(),
        "uploaded_by": cu.name,
        "source_filename": file.filename,
        "employee_count": len(employees),
        "employees": employees,
        "totals": totals,
        "company_name": company_name,
        "company_address": company_address,
    }
    # One Pay Register per period, full stop — re-uploading the same month overwrites
    # whatever was there before rather than piling up versions with no UI to see them.
    db[collections.PAYROLL_REGISTERS].replace_one(
        {"org_id": cu.org_id, "period": detected_period}, doc, upsert=True,
    )
    log_action(
        db, user=cu.name, action="CREATE", module="Payroll", org_id=cu.org_id, record_id=detected_period,
        detail=f"Uploaded Pay Register for {detected_period} ({len(employees)} employees)",
    )
    out = _out(doc)
    out["employees"] = _flag_unmatched_employees(out["employees"], db, cu.org_id)
    return out


@router.get("/periods", dependencies=[Depends(require_permission("Payroll", "view"))])
def list_periods(db: Database = Depends(get_db), cu=Depends(get_current_user)):
    rows = (
        db[collections.PAYROLL_REGISTERS]
        .find({"org_id": cu.org_id}, {"employees": 0, "totals": 0})
        .sort("period", -1)
    )
    return [_out_meta(d) for d in rows]


def _list_employee_fields(db: Database, org_id: str, keys):
    rows = db[collections.EMPLOYEES].find({"org_id": org_id}).sort("name", 1)
    return [_out_employee_fields(e, keys) for e in rows]


def _update_employee_fields(db: Database, org_id: str, emp_id: str, patch: dict, keys, detail_label: str, cu):
    e = get_or_404(db, collections.EMPLOYEES, emp_id, org_id, "Employee not found")
    if patch:
        db[collections.EMPLOYEES].update_one({"_id": emp_id, "org_id": org_id}, {"$set": patch})
        e = db[collections.EMPLOYEES].find_one({"_id": emp_id, "org_id": org_id})
    log_action(db, user=cu.name, action="UPDATE", module="Payroll", org_id=org_id, record_id=emp_id,
               detail=f"Updated {detail_label} for {e['name']}")
    return _out_employee_fields(e, keys)


@router.get("/bank-details", dependencies=[Depends(require_permission("Payroll", "view"))])
def list_bank_details(db: Database = Depends(get_db), cu=Depends(get_current_user)):
    return _list_employee_fields(db, cu.org_id, _BANK_STATUTORY_KEYS)


@router.patch("/bank-details/{emp_id}", dependencies=[Depends(require_permission("Payroll", "edit"))])
def update_bank_details(emp_id: str, payload: BankStatutoryUpdate, db: Database = Depends(get_db), cu=Depends(get_current_user)):
    return _update_employee_fields(db, cu.org_id, emp_id, payload.dict(exclude_none=True), _BANK_STATUTORY_KEYS, "bank & statutory details", cu)


@router.get("/job-details", dependencies=[Depends(require_permission("Payroll", "view"))])
def list_job_details(db: Database = Depends(get_db), cu=Depends(get_current_user)):
    return _list_employee_fields(db, cu.org_id, _JOB_DETAILS_KEYS)


@router.patch("/job-details/{emp_id}", dependencies=[Depends(require_permission("Payroll", "edit"))])
def update_job_details(emp_id: str, payload: JobDetailsUpdate, db: Database = Depends(get_db), cu=Depends(get_current_user)):
    return _update_employee_fields(db, cu.org_id, emp_id, payload.dict(exclude_none=True), _JOB_DETAILS_KEYS, "job details", cu)


@router.get("/employee-details", dependencies=[Depends(require_permission("Payroll", "view"))])
def list_employee_details(db: Database = Depends(get_db), cu=Depends(get_current_user)):
    return _list_employee_fields(db, cu.org_id, _EMPLOYEE_DETAILS_EXTRA_KEYS)


@router.patch("/employee-details/{emp_id}", dependencies=[Depends(require_permission("Payroll", "edit"))])
def update_employee_details(emp_id: str, payload: EmployeeDetailsUpdate, db: Database = Depends(get_db), cu=Depends(get_current_user)):
    return _update_employee_fields(db, cu.org_id, emp_id, payload.dict(exclude_none=True), _EMPLOYEE_DETAILS_EXTRA_KEYS, "employee details", cu)


@router.patch("/{period}/company-info", dependencies=[Depends(require_permission("Payroll", "edit"))])
def update_company_info(period: str, payload: CompanyInfoUpdate, db: Database = Depends(get_db), cu=Depends(get_current_user)):
    """Manual fallback for when a Pay Register's source file didn't carry a company
    name/address in its title block (see _detect_company_info) — used on the Payslip
    header, so this lets it be filled in by hand instead."""
    doc = db[collections.PAYROLL_REGISTERS].find_one({"org_id": cu.org_id, "period": period})
    if not doc:
        raise HTTPException(404, "No Pay Register found for this period")
    patch = payload.dict(exclude_none=True)
    if patch:
        db[collections.PAYROLL_REGISTERS].update_one({"org_id": cu.org_id, "period": period}, {"$set": patch})
        doc = db[collections.PAYROLL_REGISTERS].find_one({"org_id": cu.org_id, "period": period})
    log_action(db, user=cu.name, action="UPDATE", module="Payroll", org_id=cu.org_id, record_id=period,
               detail=f"Updated company info for Pay Register {period}")
    return _out_meta(doc)


def _flag_unmatched_employees(employees: list, db: Database, org_id: str) -> list:
    """Marks each row with whether its Code matches an actual Employee record — computed
    fresh on every read (never persisted on the stored document) so a Code uploaded before
    its Employee record existed stops being flagged the moment that record is created,
    with no need to re-upload."""
    valid_codes = {e["emp_id"] for e in db[collections.EMPLOYEES].find({"org_id": org_id}, {"emp_id": 1})}
    for row in employees:
        row["employee_found"] = row.get("code") in valid_codes
    return employees


@router.get("/{period}", dependencies=[Depends(require_permission("Payroll", "view"))])
def get_register(period: str, db: Database = Depends(get_db), cu=Depends(get_current_user)):
    doc = db[collections.PAYROLL_REGISTERS].find_one({"org_id": cu.org_id, "period": period})
    if not doc:
        raise HTTPException(404, "No Pay Register found for this period")
    out = _out(doc)
    out["employees"] = _flag_unmatched_employees(out["employees"], db, cu.org_id)
    return out


@router.delete("/{period}", dependencies=[Depends(require_permission("Payroll", "delete"))])
def delete_register(period: str, db: Database = Depends(get_db), cu=Depends(get_current_user)):
    result = db[collections.PAYROLL_REGISTERS].delete_one({"org_id": cu.org_id, "period": period})
    if result.deleted_count == 0:
        raise HTTPException(404, "No Pay Register found for this period")
    log_action(db, user=cu.name, action="DELETE", module="Payroll", org_id=cu.org_id, record_id=period,
               detail=f"Deleted Pay Register for {period}")
    return {"message": "Deleted"}
